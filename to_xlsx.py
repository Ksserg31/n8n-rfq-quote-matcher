import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

res = json.load(open("tests/results.json", encoding="utf-8"))

HDR  = PatternFill("solid", fgColor="1F3B57")
WARN = PatternFill("solid", fgColor="FDF0D5")
OKF  = PatternFill("solid", fgColor="EDF5EF")
thin = Side(style="thin", color="C9D2DA")
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook(); wb.remove(wb.active)

for f in res:
    name = f["file"].replace(".pdf", "")[:28]
    ws = wb.create_sheet(f"Q_{name}"[:31])
    cols = ["Line","RFQ code","RFQ description","Qty","Unit",
            "Matched code","Catalogue description","Unit price EUR","Line total EUR",
            "Confidence","Status","Note / why"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(vertical="center", wrap_text=True); c.border = BOX
    for l in f["lines"]:
        ws.append([l["line"], l["raw_code"] or "", (l["raw_description"] or "")[:90],
                   l["qty"], l["raw_unit"] or "",
                   l["matched_code"] or "", l["matched_description"] or "",
                   l["unit_price_eur"], l["line_total_eur"],
                   l["confidence"] or "", "REVIEW" if l["status"]=="review" else "matched",
                   l["reason"] or ""])
        row = ws[ws.max_row]
        for c in row:
            c.border = BOX; c.alignment = Alignment(vertical="top", wrap_text=True)
            c.fill = WARN if l["status"]=="review" else OKF
        for i in (8,9):
            row[i-1].number_format = '#,##0.00'
    s = f["summary"]
    ws.append([])
    ws.append(["", "", "Lines", s["total_lines"], "", "Matched", s["matched"], "", "", "", "Review", s["review"]])
    ws.append(["", "", "", "", "", "", "Subtotal of matched lines", "", s["quote_subtotal_eur"]])
    ws[ws.max_row][8].number_format = '#,##0.00'
    ws[ws.max_row][8].font = Font(bold=True)
    ws.append(["", "", s["note"]])
    widths = [6,20,40,9,7,20,40,13,13,11,10,52]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

# review queue across all files
ws = wb.create_sheet("Review queue", 0)
ws.append(["File","Line","RFQ code","RFQ description","Qty","Why it stopped here",
           "Candidate 1","score","Candidate 2","score","Candidate 3","score"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR; c.border = BOX
    c.alignment = Alignment(vertical="center", wrap_text=True)
for f in res:
    for l in f["lines"]:
        if l["status"] != "review": continue
        cands = l.get("candidates") or []
        row = [f["file"], l["line"], l["raw_code"] or "", (l["raw_description"] or "")[:90],
               l["qty"], l["reason"] or ""]
        for i in range(3):
            row += [cands[i]["code"] if i < len(cands) else "",
                    cands[i]["desc_score"] if i < len(cands) else ""]
        ws.append(row)
        for c in ws[ws.max_row]:
            c.border = BOX; c.fill = WARN; c.alignment = Alignment(vertical="top", wrap_text=True)
for i,w in enumerate([26,6,18,42,8,54,18,8,18,8,18,8],1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"

wb.save("samples/quotation_draft.xlsx")
print("samples/quotation_draft.xlsx")
